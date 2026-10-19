"""Regenerates docs/Cropin_API_Test_Scenarios.xlsx from the ROWS list below —
the single source of truth for the test-case catalogue, kept as a script
(not hand-edited in Excel) so re-running it always produces the same
formatting. ROWS is meant to be extended, not the workbook.

Whenever a new test function is scaffolded (per the intake workflow in
CLAUDE.md / README.md section 6), append a matching tuple to ROWS — same
five columns, same order the test's own module section already uses — then
re-run this script:

    python scripts/build_test_scenarios_xlsx.py

Test Case ID convention: TC-<MODULE>-NNN, numbered sequentially within its
module section (e.g. the next farmer test after TC-FARMER-002 is
TC-FARMER-003). A brand-new module gets its own MODULE_SECTION_STARTS entry
and starts at NNN=001.

Test Case Name convention: a short human-readable title, e.g. "Farmer
Creation" or "Farmer Edit" — not the raw pytest function name. Put the
actual function (or cleanup step) name as a `backticked` prefix at the
start of the "What It Will Do" text instead, e.g.
"(`test_create_farmer_success`) Creates the one shared farmer record...",
so it's still easy to find the code from a row without it being the
column someone reads first.
"""
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_ROOT, "docs", "Cropin_API_Test_Scenarios.xlsx")

HEADERS = ["Test Case ID", "Test Case Name", "Test Type", "What It Will Do", "APIs Used"]

# Test Type reflects the actual pytest markers on each test function (see
# tests/*/test_*.py) — most business-flow tests carry both sanity and
# regression, but the auth negative-path tests only carry regression
# (they're deliberately excluded from sanity runs), the framework
# self-tests carry no test-type marker at all, and the CLEANUP section's
# rows are teardown steps rather than pytest tests — tagged "Cleanup"
# since they don't carry a sanity/regression marker at all, but run every
# time the suite does (see tests/conftest.py's cleanup_created_records).
# "smoke" is registered in pyproject.toml but currently applied to no
# test — removed from every test for now (still easy to bring back).
# Test Case Name is a human-readable title (e.g. "Farmer Creation") — the
# underlying pytest function / cleanup step name is kept as a `backticked`
# prefix on the "What It Will Do" text instead, so nothing is lost for
# grepping the code, it's just not the headline column anymore.
ROWS = [
    # --- AUTH ---
    ("TC-AUTH-001", "Login - Valid Credentials", "Sanity, Regression",
     "(`test_generate_token_success`) Happy-path login: requests an access token using valid tenant/credentials and verifies the response is 200 with a valid token schema. Foundational test — if it fails, the entire run aborts instead of just this test failing.",
     "POST /auth/realms/{tenant_code}/protocol/openid-connect/token"),
    ("TC-AUTH-002", "Login - Token Persisted to Session", "Sanity, Regression",
     "(`test_generate_token_stores_tokens_on_session`) Verifies that after a successful login, the access_token and refresh_token are correctly stored on the shared session context so every other service call can reuse them.",
     "POST /auth/realms/{tenant_code}/protocol/openid-connect/token"),
    ("TC-AUTH-003", "Login - Invalid Credentials", "Regression",
     "(`test_generate_token_fails_with_invalid_credentials`) Negative: attempts login with an invalid username/password and verifies the API rejects it with 401.",
     "POST /auth/realms/{tenant_code}/protocol/openid-connect/token"),
    ("TC-AUTH-004", "Login - Invalid Tenant Code", "Regression",
     "(`test_generate_token_fails_with_invalid_tenant_code`) Negative: attempts login against an unknown/invalid tenant realm and verifies the API returns 400 or 404.",
     "POST /auth/realms/{invalid_tenant_code}/protocol/openid-connect/token"),
    ("TC-AUTH-005", "Login - Invalid Client Secret", "Regression",
     "(`test_generate_token_fails_with_invalid_client_secret`) Negative: attempts login with an invalid client_secret and verifies the API rejects it with 400.",
     "POST /auth/realms/{tenant_code}/protocol/openid-connect/token"),
    ("TC-AUTH-006", "Login - Missing Grant Type", "Regression",
     "(`test_generate_token_fails_when_grant_type_missing`) Negative: sends a token request with the grant_type field missing and verifies the API rejects it with 400.",
     "POST /auth/realms/{tenant_code}/protocol/openid-connect/token"),

    # --- FARMER ---
    ("TC-FARMER-001", "Farmer Creation", "Sanity, Regression",
     "(`test_create_farmer_success`) Creates the one shared farmer record used for the whole run and verifies the create call returns 201.",
     "POST /services/farm/api/farmers"),
    ("TC-FARMER-002", "Farmer Edit", "Sanity, Regression",
     "(`test_edit_farmer_updates_first_name`) Fetches the shared farmer, appends a timestamp to firstName, PUTs the updated record back, and verifies the update returns 200.",
     "GET /services/farm/api/farmers/{id}\nPUT /services/farm/api/farmers"),

    # --- ASSET ---
    ("TC-ASSET-001", "Asset Creation", "Sanity, Regression",
     "(`test_create_asset_success`) Creates the one shared asset (owned by the shared farmer) used for the whole run and verifies the create call returns 201.",
     "POST /services/farm/api/assets"),
    ("TC-ASSET-002", "Asset Edit", "Sanity, Regression",
     "(`test_edit_asset_updates_name`) Fetches the shared asset, appends a timestamp to name, PUTs the updated record back, and verifies the update returns 200.",
     "GET /services/farm/api/assets/{id}\nPUT /services/farm/api/assets"),

    # --- PLAN TYPE ---
    ("TC-PLANTYPE-001", "Plan Type Creation", "Sanity, Regression",
     "(`test_create_plantype_success`) Creates the one shared plan type used for the whole run and verifies the create call returns 201.",
     "POST /services/farm/api/plan-types"),
    ("TC-PLANTYPE-002", "Plan Type Edit", "Sanity, Regression",
     "(`test_edit_plantype_updates_name`) Fetches the shared plan type, appends a timestamp to name, PUTs the updated record back, and verifies the update returns 200.",
     "GET /services/farm/api/plan-types/{id}\nPUT /services/farm/api/plan-types"),

    # --- CROP VARIETY ---
    ("TC-VARIETY-001", "Crop Variety Creation", "Sanity, Regression",
     "(`test_create_crop_variety_success`) Resolves crop-stage and harvest-grade ids live by name, creates the one shared crop variety, then adds its mandatory additional-attribute record; verifies both calls return 201.",
     "GET /services/farm/api/crop-stages\nGET /services/master/api/harvest-grades\nPOST /services/farm/api/varieties\nPUT /services/farm/api/additional-attribute"),
    ("TC-VARIETY-002", "Crop Variety Edit", "Sanity, Regression",
     "(`test_edit_crop_variety_updates_name`) Fetches the shared crop variety, appends a timestamp to name, PUTs the updated record back, and verifies the update returns 200.",
     "GET /services/farm/api/varieties/{id}\nPUT /services/farm/api/varieties"),

    # --- PLAN ---
    ("TC-PLAN-001", "Plan Creation (Linked to Crop Variety)", "Sanity, Regression",
     "(`test_add_plan_to_crop_variety_success`) Links the shared plan type and crop variety by creating a plan — reads the plan type's current details live as part of building the payload — and verifies the create call returns 200.",
     "GET /services/farm/api/plan-types/{id}\nPOST /services/farm/api/plans"),

    # --- PROJECT ---
    ("TC-PROJECT-001", "Project Creation", "Sanity, Regression",
     "(`test_create_project_success`) Creates the one shared project used for the whole run and verifies the create call returns 201.",
     "POST /services/farm/api/projects"),
    ("TC-PROJECT-002", "Add Asset to Project", "Sanity, Regression",
     "(`test_add_asset_to_project_success`) Adds the shared asset to the shared project as a probable asset and verifies at least one record completed successfully.",
     "POST /services/farm/api/projects/{id}/probable-assets"),
    ("TC-PROJECT-003", "Validate Project Assets", "Sanity, Regression",
     "(`test_validate_project_assets_success`) Self-validates the project-asset association and verifies it returns non-empty croppableAreaIds with zero failed records.",
     "POST /services/farm/api/projects/{id}/self-validate-project-assets?cloneFlag=false"),
    ("TC-PROJECT-004", "Project Execution", "Sanity, Regression",
     "(`test_execute_project_completes`) Starts execution on the shared project and polls the async job until it finishes, verifying the final status is COMPLETED.",
     "POST /services/farm/api/projects/{id}/start-execution\nGET /services/farm/api/bulk-process/{execution_id} (polled)"),
    ("TC-PROJECT-005", "Croppable Area - Add Variety and Sowing Date", "Sanity, Regression",
     "(`test_add_variety_and_sowing_date_to_croppable_area`) Fetches the croppable area auto-created for the project, updates it with the shared crop variety and a sowing date, and verifies the PUT response echoes both back correctly.",
     "GET /services/farm/api/croppable-areas/{id}\nPUT /services/farm/api/croppable-areas"),
    ("TC-PROJECT-006", "Croppable Area - Verify Creation", "Sanity, Regression",
     "(`test_verify_croppable_area_created`) Re-fetches the croppable area by id and verifies the returned id matches the one produced during asset validation.",
     "GET /services/farm/api/croppable-areas/{id}"),
    ("TC-PROJECT-007", "Croppable Area - Verify Variety and Sowing Date", "Sanity, Regression",
     "(`test_verify_croppable_area_variety_and_sowing_date`) Re-fetches the croppable area with a fresh GET (not the PUT echo) to confirm the variety and sowing date update actually persisted on the server.",
     "GET /services/farm/api/croppable-areas/{id}"),
    ("TC-PROJECT-008", "Croppable Area - Verify Plan Type Tasks", "Sanity, Regression",
     "(`test_verify_plan_type_tasks_pulled_to_croppable_area`) Verifies that once the variety (linked to a plan and plan type) is assigned to the croppable area, the plan's tasks — with the correct planId, plan type name, and scheduled dates — pull through onto the croppable area.",
     "GET /services/farm/api/tasks/croppablearea/{id}?sort=lastModifiedDate,desc"),

    # --- CLEANUP (teardown steps in tests/conftest.py's cleanup_created_records
    # autouse fixture, not standalone pytest test functions — no pass/fail row of
    # their own, run once at the end of every run in this fixed order, each
    # skipped individually if its own entity was never created this run) ---
    ("TC-CLEANUP-001", "Cleanup - Close and Remove Croppable Area", "Cleanup",
     "(`cleanup_close_and_remove_croppable_area`) Closes the croppable area(s) created via project-asset validation, then removes them through the project's project-assets endpoint (an async job, polled to completion). First step of teardown; skipped if the project itself was never created this run.",
     "GET /services/farm/api/croppable-areas/closed?reasonId=4&ids=...\nPOST /services/farm/api/projects/{id}/project-assets/remove/selected-ids\nGET /services/farm/api/bulk-process/{job_id} (polled)"),
    ("TC-CLEANUP-002", "Cleanup - Delete Plan", "Cleanup",
     "(`cleanup_delete_plan`) Deletes the shared plan created this run. Runs before the crop variety delete below, since the plan references that variety.",
     "DELETE /services/farm/api/plans/bulk?ids=..."),
    ("TC-CLEANUP-003", "Cleanup - Delete Crop Variety", "Cleanup",
     "(`cleanup_delete_crop_variety`) Deletes the shared crop variety created this run.",
     "DELETE /services/farm/api/varieties/bulk?ids=..."),
    ("TC-CLEANUP-004", "Cleanup - Delete Asset", "Cleanup",
     "(`cleanup_delete_asset`) Deletes the shared asset created this run.",
     "DELETE /services/farm/api/assets/bulk?ids=..."),
    ("TC-CLEANUP-005", "Cleanup - Delete Farmer", "Cleanup",
     "(`cleanup_delete_farmer`) Deletes the shared farmer created this run.",
     "DELETE /services/farm/api/farmers/bulk?ids=..."),
    ("TC-CLEANUP-006", "Cleanup - Stop Project Execution", "Cleanup",
     "(`cleanup_stop_project_execution`) Stops the shared project's execution — an async job, polled to completion — required before the project's other records can be considered fully torn down.",
     "POST /services/farm/api/projects/{id}/stop-execution\nGET /services/farm/api/bulk-process/{job_id} (polled)"),
    ("TC-CLEANUP-007", "Cleanup - Delete Plan Type", "Cleanup",
     "(`cleanup_delete_plan_type`) Deletes the shared plan type created this run. Last step of the whole teardown chain, run only once every other record above is already gone.",
     "DELETE /services/farm/api/plan-types/bulk?ids=..."),

    # --- FRAMEWORK (intake validation gate — internal unit tests, no live API calls) ---
    ("TC-FW-001", "Intake Spec - Valid Spec Accepted", "N/A (no marker)",
     "(`test_valid_spec_passes`) Verifies a fully valid single-scenario intake spec passes schema validation.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-002", "Intake Spec - Unknown Module Rejected", "N/A (no marker)",
     "(`test_unknown_module_rejected`) Verifies an intake spec naming a module outside the allowed set (farmer/asset/plantype/crop_variety/project/login) is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-003", "Intake Spec - Unsupported Method Rejected", "N/A (no marker)",
     "(`test_unsupported_method_rejected`) Verifies an intake spec using an HTTP method outside the allowed set is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-004", "Intake Spec - Unknown Test Type Rejected", "N/A (no marker)",
     "(`test_unknown_test_type_rejected`) Verifies an intake spec with a test_type outside sanity/smoke/regression is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-005", "Intake Spec - Endpoint Must Be Relative Path", "N/A (no marker)",
     "(`test_endpoint_must_be_relative_path`) Verifies an intake spec whose endpoint doesn't start with \"/\" is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-006", "Intake Spec - Body Method Without Payload Rejected", "N/A (no marker)",
     "(`test_body_method_without_payload_rejected`) Verifies a POST/PUT/PATCH intake spec with an empty payload is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-007", "Intake Spec - GET Without Payload Accepted", "N/A (no marker)",
     "(`test_get_without_payload_is_fine`) Verifies a GET intake spec with no payload is accepted, since payload is optional for GET.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-008", "Intake Spec - Negative Scenario Requires Error Status", "N/A (no marker)",
     "(`test_negative_scenario_requires_error_status`) Verifies an intake spec marked negative:true but with a non-error expected_status is rejected.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-009", "Intake Spec - Negative Scenario With Error Status Accepted", "N/A (no marker)",
     "(`test_negative_scenario_with_error_status_passes`) Verifies an intake spec marked negative:true with a proper >=400 expected_status is accepted.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-010", "Intake Spec - Batch Validates All Items", "N/A (no marker)",
     "(`test_batch_validates_all_items`) Verifies a batch of several valid intake specs all validate together successfully.",
     "N/A (unit test, no live API call)"),
    ("TC-FW-011", "Intake Spec - Batch Rejects Invalid Item", "N/A (no marker)",
     "(`test_batch_rejects_if_any_item_invalid`) Verifies that if one scenario in an intake batch is invalid, validation of the whole batch fails.",
     "N/A (unit test, no live API call)"),
]

MODULE_SECTION_STARTS = {
    0: "AUTH",
    6: "FARMER",
    8: "ASSET",
    10: "PLAN TYPE",
    12: "CROP VARIETY",
    14: "PLAN",
    15: "PROJECT",
    23: "CLEANUP",
    30: "FRAMEWORK (Intake Validation)",
}

NUM_COLS = len(HEADERS)

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
SECTION_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1F4E78")
BODY_FONT = Font(name=FONT_NAME, size=10)
ID_FONT = Font(name=FONT_NAME, size=10, bold=True)
thin = Side(style="thin", color="B7B7B7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_TOP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
CENTER_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test Scenarios"

# Title row
last_col_letter = chr(ord("A") + NUM_COLS - 1)
ws.merge_cells(f"A1:{last_col_letter}1")
ws["A1"] = "Cropin API Automation — Test Scenario Catalogue"
ws["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 26

ws.merge_cells(f"A2:{last_col_letter}2")
ws["A2"] = "41 test scenarios across auth/farmer/asset/plantype/crop_variety/plan/project, the cleanup teardown chain, and the intake-validation framework self-tests"
ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color="595959")
ws.row_dimensions[2].height = 16

HEADER_ROW = 4

for col_idx, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER_TOP
    cell.border = BORDER
ws.row_dimensions[HEADER_ROW].height = 20

current_row = HEADER_ROW + 1
for idx, row_data in enumerate(ROWS):
    if idx in MODULE_SECTION_STARTS:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=NUM_COLS)
        sec_cell = ws.cell(row=current_row, column=1, value=MODULE_SECTION_STARTS[idx])
        sec_cell.font = SECTION_FONT
        sec_cell.fill = SECTION_FILL
        sec_cell.alignment = Alignment(horizontal="left", vertical="center")
        for c in range(1, NUM_COLS + 1):
            ws.cell(row=current_row, column=c).border = BORDER
        ws.row_dimensions[current_row].height = 18
        current_row += 1

    tc_id, tc_name, tc_type, tc_desc, tc_apis = row_data
    ws.cell(row=current_row, column=1, value=tc_id).font = ID_FONT
    ws.cell(row=current_row, column=2, value=tc_name).font = BODY_FONT
    ws.cell(row=current_row, column=3, value=tc_type).font = BODY_FONT
    ws.cell(row=current_row, column=4, value=tc_desc).font = BODY_FONT
    ws.cell(row=current_row, column=5, value=tc_apis).font = BODY_FONT

    for c in range(1, NUM_COLS + 1):
        cell = ws.cell(row=current_row, column=c)
        cell.alignment = WRAP_TOP_CENTER if c == 3 else WRAP_TOP_LEFT
        cell.border = BORDER

    current_row += 1

# Column widths
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 42
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 68
ws.column_dimensions["E"].width = 52

# Freeze header
ws.freeze_panes = f"A{HEADER_ROW + 1}"

wb.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Total data rows:", len(ROWS))
